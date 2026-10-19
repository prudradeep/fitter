from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "process_flow.xlsx"


def sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174A6E")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for idx, width in enumerate(widths or [24] * len(headers), 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 30
    return ws


wb = Workbook()
ws = wb.active
ws.title = "Read Me"
ws.append(["Dr Transition workflow map", "Implementation-based guide"])
ws.append(["Scope", "From initial Country selection through Region, Sector, Policy, hazard and population context, mitigation choice, proposal/adaptation branches, shared review, summary and evaluation."])
ws.append(["How to read", "Each row is a step or decision. Branch rows identify their own next journey. Option sources distinguish database/reference data, user input, runtime retrieval, LLM-generated content and fixed interface actions."])
ws.append(["Coverage note", "The initial Country → Region → Sector selection is database-backed. Policy selection is policy/reference data plus user-added policy paths. Suggested content later in the flow is source-grounded runtime output, not a fixed list of options. Exact prompts and available choices can vary with configured data and retrieved evidence."])
ws.append(["Primary references", "app/services/chat_navigation_steps.py; app/services/chat_hazard_steps.py; app/services/chat_mitigation_steps.py; app/services/chat_mitigation_creation_guided.py; app/seed/sqlite_basic_data.py; mm_flow.md; mitigation_options_flow.md"])
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 125
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", size=14)
    c.fill = PatternFill("solid", fgColor="174A6E")
for row in ws.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
ws.sheet_view.showGridLines = False

sheet(wb, "End to End Flow",
      ["#", "Step / decision", "Description", "Options / branch and next journey", "Option or content source", "Validation / transition logic", "Implementation / reference"], [
 [1,"Select Country","Set the country context for the session.","Choose one available country → Region.","DB: countries table; seeded bootstrap/reference rows. Current SQLite seed includes Germany, Hungary, Ireland, Italy, Spain, Portugal.","Selection must match an available country; session country is set and dependent context follows it.","app/seed/sqlite_basic_data.py; app/services/chat_navigation_steps.py"],
 [2,"Select Region","Choose a region within the selected country.","Choose listed region → Sector.","DB: regions filtered by country_id.","Only regions attached to selected country are offered/accepted.","app/models.py (Region); app/services/chat_navigation_steps.py"],
 [3,"Select Sector","Choose an analytical sector supported for the selected country.","Energy / Housing / Transport where configured → Policy/context and hazard analysis. Not every country has every sector.","DB: sectors joined through country_sectors; SQLite bootstrap seeds configured country-sector pairs.","Options constrained to selected country's configured sector links.","app/models.py (Sector, country_sectors); app/seed/sqlite_basic_data.py"],
 [4,"Select / review Policy","Select relevant policy and review its supplied policy context; user-added policy can be available through the flow.","Choose existing listed policy → policy context; add/provide a policy where offered → user-policy path; policy summary may ask for clarification or source/reference input.","DB/reference policy records and imported policy data; user-provided policy details/documents; linked policy documents and knowledge base for context.","Policy facts are grounded in selected summary, references and documents. Missing summary fields can trigger reference/clarification; unsupported assertions are not treated as established facts.","app/services/chat_hazard_steps.py; app/seed/reference_data.py; mm_flow.md"],
 [5,"Review hazards","View/rank hazards relevant to selected country, region, sector and data; choose one, or enter custom hazard through available path.","Listed hazard → selected-hazard journey; expert/additional hazards → expert catalog; co-created/custom hazard → custom hazard acceptance and then same downstream journey; add hazard → validation journey.","DB/reference system hazards and additional hazards; ranked/API results; custom user text and generated/validated profile data.","Hazard ranking uses context; custom entries pass quality, relevance, duplicate, evidence/grounding and acceptance checks as applicable. Invalid/unclear inputs loop for correction.","app/services/chat_hazard_steps.py; app/services/chat_hazard_catalog.py; app/services/chat_custom_hazard_creation.py"],
 [6,"Select affected groups / profiles","Identify socio-demographic groups affected by the selected hazard.","Select available profile(s), answer configured questions, or use relevant custom hazard population flow → review.","DB-configured evaluation questions/options and profile/hazard records; in custom paths, user input plus grounded model extraction.","Collect selections/answers; profile review can allow correction; custom groups require clarity/relevance and duplicate/grounding checks.","app/seed/sqlite_basic_data.py; app/services/chat_hazard_steps.py; app/services/chat_custom_hazard_population_steps.py"],
 [7,"Review population and choose mitigation route","Confirm affected groups and decide whether to start mitigation planning.","Continue analysis / other available navigation → corresponding step; Start Mitigation Planning → choose a hazard then mitigation entry; other hazard-list options return to hazard listing.","Fixed workflow actions plus selected hazard/profile context.","Must have hazard context for mitigation; route preserves accepted custom or catalog hazard and population context.","mm_flow.md; app/services/chat_hazard_steps.py; app/services/chat_navigation_steps.py"],
 [8,"Choose mitigation entry","Open guided mitigation planning for the selected hazard.","Choose mechanism from suggestions or enter one → find supporting policy; question/other-options navigation stays in or returns to the appropriate stage.","Mechanisms and general suggestions are composed at runtime from hazard context and knowledge/reference material; free-text option is user input.","Mechanism must be selectable or sufficiently clear; downstream policy search is tied to chosen mechanism and hazard.","mm_flow.md; app/services/chat_mitigation_creation_guided.py"],
 [9,"Find and confirm supporting policy","Search for a policy supporting the selected mechanism-to-hazard relationship.","Supported result → inspect policy linkage and confirm; reject/no supported result → provide policy URL/file → validate/clarify/upload again until supported.","Runtime retrieval: linked custom-hazard references, core and validated-evidence knowledge base, policy materials; supplied URL/file is user input.","Check policy relevance and causal linkage to chosen mechanism/hazard. Unsupported or unclear result loops to clarification or evidence input.","mm_flow.md; app/services/chat_mitigation_creation_guided.py"],
 [10,"Choose proposal journey","After policy/mechanism context, choose how to develop the mitigation.","Create a new mitigation proposal → Steps 11A onward. Propose adaptations in existing policy → Steps 11B onward. Adopt/discard inspiration is a later optional branch.","Two fixed workflow options (new_policy / existing_policy); suggestions/inspirations are runtime source-backed content.","Persist branch in session and carry selected country, region, sector, hazard, policy and groups forward.","mitigation_options_flow.md; app/services/chat_mitigation_creation_guided.py"],
 ["11A","New proposal: type, title, description","Design a new policy action; choose one or more proposal types and enter proposal details.","Regulation / Economic incentives / Public services / Information/education; then enter title and description. No recognized type → reselect. Revise review → edit proposal.","Type labels are fixed application options; title/description are user-authored.","At least one recognized type. Preliminary review assesses specific action, objectives, sector objective, hazard relevance, pathway and material factual claims; unclear action/pathway prompts clarification. Missing evidence support may offer evidence or continue without it.","mitigation_options_flow.md; app/services/chat_mitigation_creation_guided.py"],
 ["11B","Existing policy: context and gap","Anchor the interview to a selected policy and explain the policy gap (Question 1).","Confirm policy or select one if absent; inspect mechanisms/challenges/adaptation suggestions; answer gap → confirm/revise reflection.","Selected policy/session data; matching D4.2 factsheet, linked policy-document passages and shared KB; user-added policy may use same-sector semantic factsheet examples.","Ground mechanisms/adaptations in source records. Future proposed adjustment is not represented as current policy. No supported source → state none / request evidence where applicable.","mitigation_options_flow.md; app/services/chat_mitigation_creation_guided.py; app/services/policy_summary_references.py"],
 ["12","Existing policy: propose adjustment (Q2)","Describe a concrete change that responds to the validated gap.","Clear and linked → target groups; unclear → clarify/revise Q2.","User-authored adjustment; checked against selected policy, gap, hazard and source context.","Must respond to Q1, plausibly mitigate selected hazard and connect to an actual policy mechanism. Proposed future change is not treated as a current provision.","mitigation_options_flow.md"],
 ["13","Target population and group pathway","Select groups and explain how each group benefits from the proposal/adaptation.","Select one or more configured demographic options per question or skip; explain group-specific barrier/benefit; clarify or optionally provide evidence if unclear/unsupported.","Configured questions/options from DB; group pathway is user-authored and validated with selected context and available sources.","Check target group, clarity, relevance to measure/policy/hazard and support for factual claims. Missing group/pathway can return to input/review; groups can be added/removed.","app/seed/sqlite_basic_data.py; app/services/chat_mitigation_creation_guided.py; mitigation_options_flow.md"],
 ["14","Shared mechanism reflection","Explain how the measure changes selected mechanism and reduces hazard.","Confirm reflection → effects review; clarify/revise → re-evaluate.","Runtime generated reflection from proposal/measure, hazard, mechanism, policy linkage, groups, pathways and sources; user confirmation/correction.","Check specificity and fit to mechanism/hazard; loop until confirmed or sufficiently clarified.","mm_flow.md; app/services/chat_mitigation_creation_guided.py"],
 ["15","Policy effects and inspiration","Review grounded unintended effects on other policy aspects; inspect related FITTER Open Labs proposals/adjustments.","For each grounded effect: Yes → clarify and record additional mitigation; No → record reason. Use inspiration fully / adopt parts / discard; adopted changes re-run reflection, evidence, effects and summary checks.","Effect candidates and inspirations retrieved from policy/hazard records and knowledge sources; decision is user input.","Only grounded distinct effects are presented; if none, skip. Adoption re-validates affected content and returns to checks when support is missing.","mm_flow.md; mitigation_options_flow.md"],
 ["16","Evidence decision and validation","Supply optional evidence for measure and mechanism where useful.","Yes + URL/file/text → relevance validation; No/skip → continue without supplied evidence; invalid/irrelevant → clarify or re-enter.","User evidence, URL/file, extracted document content; retrieved source passages remain separate source context.","Evidence must be relevant to both measure and mechanism. Evidence is optional; unsupported status is retained in the summary rather than silently asserted as verified.","mm_flow.md; app/services/chat_mitigation_creation_workflow.py"],
 ["17","Final group review and summary","Review all selected groups, pathways, measure, mechanism, evidence and policy-effect decisions.","Confirm summary → save/review; revise → return to the relevant input stage and rerun checks.","Structured session state plus confirmed user answers and available source excerpts; summary content is assembled by app, with generated/reflected text where noted.","Summary carries proposal/measure, mechanisms, confirmed reflection, evidence status, policy effects, groups and group-evidence status. Missing support can route back to clarification; user confirms or modifies.","mitigation_options_flow.md; mm_flow.md; app/services/chat_mitigation_creation_guided.py"],
 [18,"Save, evaluate, complete","Persist the mitigation and continue to concept comparison/evaluation/system inquiry as configured.","Save → mitigation review/recorded → concept comparison → evaluation → system inquiry/completion.","DB persistence of accepted session records; configured evaluation questions and user answers.","Final checks gate saving/review; evaluation choices and scores are collected against configured questions. Exact end steps depend on selected route/configuration.","mm_flow.md; app/services/chat_mitigation_steps.py; app/seed/sqlite_basic_data.py"],
])

sheet(wb, "Option Sources",
      ["Stage", "Option / content", "Source category", "Where generated / loaded", "Dependency / notes", "References"], [
 ["Country","Available countries","Database/reference","countries table seeded by reference/bootstrap seed","Current SQLite seed has six supported European countries; DB can hold runtime rows.","app/seed/sqlite_basic_data.py; app/models.py"],
 ["Region","Regions for selected country","Database","regions filtered by country_id","Dependent on selected country.","app/models.py; app/services/chat_navigation_steps.py"],
 ["Sector","Sectors for selected country","Database relation","country_sectors association joining sectors","Country-sector support varies.","app/models.py; app/seed/sqlite_basic_data.py"],
 ["Policy","Existing policies","DB/reference data","Policy records, imported reference data and policy context services","Policy can be supplemented by linked documents and KB; user-added policy has special fallback logic.","app/seed/reference_data.py; app/services/chat_hazard_steps.py"],
 ["Hazard","System/listed hazards","DB/reference + ranking","Hazard catalog and context ranking services","Context includes country/region/sector; list may include additional expert hazards.","app/services/chat_hazard_catalog.py; app/routes/api.py"],
 ["Hazard","Custom/co-created hazard","User + runtime generation/validation","User input and custom hazard service pipeline","Accepted custom hazard continues into downstream flow.","app/services/chat_custom_hazard_creation.py"],
 ["Affected profiles","Population questions and fixed answer choices","Database/configured reference","evaluation_questions and question_options seed/data","User may select multiple options or skip depending on question/path.","app/seed/sqlite_basic_data.py"],
 ["Mechanism","Suggested or entered mechanism","Runtime retrieval + user input","Retrieved hazard/policy context and knowledge base; free-text mechanism accepted for checking","Source-backed linkage required for policy search.","app/services/chat_mitigation_creation_guided.py"],
 ["Proposal type","Four proposal categories","Hard-coded application labels","new proposal branch UI/service","At least one recognized label required.","mitigation_options_flow.md"],
 ["Proposed action, gap, adjustment, reflection","User-authored plus model-generated review","User input + LLM/runtime retrieval","LLM validators/review use selected context and retrieved passages; reflection/summary uses stored session values","Outputs can prompt clarification; do not confuse generated inference with DB option labels.","app/services/chat_mitigation_creation_guided.py; mitigation_options_flow.md"],
 ["Evidence / policy suggestions","URLs, files, excerpts, supported source points","User-supplied + DB/knowledge retrieval","Upload/URL handling and knowledge/policy document retrieval","Availability varies; unsupported material leads to no-evidence/clarification branch.","app/services/chat_mitigation_creation_guided.py; mitigation_options_flow.md"],
 ["Policy-effect candidates / Open Labs inspiration","Grounded effects and related proposals","Runtime source retrieval","Policy/hazard records and knowledge sources","Can be skipped if no grounded distinct effect or inspiration is available.","mitigation_options_flow.md"],
])

sheet(wb, "Validation and Summary",
      ["Validation / summary point", "Inputs checked or used", "Logic / outcome", "Failure or alternate route", "Source"], [
 ["Country / region / sector selection","Selected values and reference rows","Country must exist; region belongs to country; sector must be linked to selected country.","Return invalid selection or present valid dependent options.","app/models.py; app/seed/sqlite_basic_data.py; app/services/chat_navigation_steps.py"],
 ["Policy summary/context","Selected policy, summary fields, linked docs, D4.2 and KB excerpts","Prefer supplied selected-policy summary/context; fallback to supported reference fields or documents. Missing facts remain unstated rather than invented.","Ask for policy source/reference or clarification; report no grounded source when unavailable.","app/services/chat_hazard_steps.py; mitigation_options_flow.md"],
 ["Hazard creation/selection","Context, user text, known hazards, evidence and sources","Catalog/ranking narrows options; custom path checks clarity, scope, similarity/duplicates and grounding/evidence before acceptance.","Clarify, revise, add evidence, reject duplicate, or return to listed hazards.","app/services/chat_hazard_steps.py; app/services/chat_custom_hazard_creation.py"],
 ["New proposal preliminary validation","Title/description/type, objectives, sector objective, hazard, mechanism, retrieved passages","Assess understandable concrete action, policy/sector objective links, hazard relevance/pathway, and material factual support. Objective-link uncertainty is advisory; unsupported material claims are reported separately.","Clarify unclear action/pathway; optionally supply evidence or continue without it; review then confirm/revise.","mitigation_options_flow.md"],
 ["Existing-policy gap and adjustment","Policy mechanisms, Q1 gap, Q2 adjustment, hazard and supporting sources","Q1 should explain how a current mechanism contributes to hazard and policy inadequacy. Q2 must answer the validated gap and link to an actual policy mechanism while reducing hazard.","Clarify/revise the answer; request evidence where configured; future adjustment cannot be presented as existing law.","mitigation_options_flow.md"],
 ["Target-group pathway","Group selection, measure/adjustment, barrier and supporting sources","Check named target group, concrete group-specific benefit/barrier, clarity and relevance to proposal/policy/hazard; unsupported factual claims may trigger optional evidence.","Clarify pathway, add/remove groups, or return to target-group entry.","mitigation_options_flow.md; app/services/chat_mitigation_creation_guided.py"],
 ["Mechanism reflection","Measure, mechanism, hazard, policy linkage, groups and evidence","Generated reflection explains how the measure changes the chosen mechanism and reduces hazard; user confirms or supplies clarification.","Revise/clarify then validate again.","mm_flow.md; app/services/chat_mitigation_creation_guided.py"],
 ["Supplied evidence","Evidence text/URL/file, measure, selected mechanism","Relevance check ties evidence to both the measure and mechanism; evidence remains optional.","Clarify, retry/upload, or continue with evidence absent/unverified status.","mm_flow.md; app/services/chat_mitigation_creation_workflow.py"],
 ["Unintended policy effects","Grounded policy/hazard effects and user decision","Only grounded distinct effects are surfaced; Yes requires explanation/additional mitigation; No requires reason.","Skip when no grounded effect exists; unresolved effects can send user back to input.","mm_flow.md; mitigation_options_flow.md"],
 ["Final summary assembly","Proposal/measure, types, mechanisms, reflection, evidence status, effects, groups/pathways and citations","Compile structured session facts and confirmed/generated content; annotate missing/unsupported support; present source controls where available.","Confirm to proceed; modify inputs to regenerate checks/summary. LLM fallback review is shown if preliminary result is unavailable or malformed.","mitigation_options_flow.md; app/services/chat_mitigation_creation_guided.py"],
 ["Persistence / evaluation","Confirmed mitigation and configured evaluation questions","Save accepted content, then proceed to concept comparison/evaluation/system inquiry; evaluation question text/options are configured data.","Return to review/input if required information is absent or validation rejects content.","mm_flow.md; app/services/chat_mitigation_steps.py; app/seed/sqlite_basic_data.py"],
])

for ws in wb.worksheets:
    if ws.max_row > 1 and ws.title != "Read Me":
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        safe = "T" + "".join(ch for ch in ws.title if ch.isalnum())[:20]
        tab = Table(displayName=safe, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(tab)

wb.save(OUT)
print(OUT)
